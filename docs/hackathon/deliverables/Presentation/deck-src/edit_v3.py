"""Round 3 on the lead's hand-edited deck (6 Oct 2026): contents slide, competitor comparison,
expected-impact slide, a clearer safety slide, renumbered judges' guide.

    python edit_v3.py <lead-edited.pptx>      -> ../Al-Sadiq-Al-Sadouq-Final.pptx

The input is the lead's own edit of the restructure.py deck, so this script edits it in place
(new slides are clones of its own template slides) instead of rebuilding from scratch.
Competitor facts are as published by Oct 2026 (sources in the comparison slide's notes).
"""
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from build_deck import NAVY, PURPLE, TURQ, OFF, BODY, MUTED, CARD, CARDLINE, MIN_PT, \
    clone_slide, ph, by_name, remove, fill, add_text, add_box, add_arrow, kicker
from restructure import icon, notes, X0, W0, OUTFILE


def base(prs, src):
    """A clone of a dark content slide with only the template's own shapes left."""
    s = clone_slide(prs, src)
    for sh in list(s.shapes):
        if not sh.name.startswith("Google Shape;"):
            remove(sh)
    return s


def build(path):
    prs = Presentation(path)
    S = list(prs.slides)
    guide, solution, safety = S[1], S[3], S[14]

    # ---- solution: the tagline sat under the lead's bullet list
    by_name(solution, "TextBox 552").top = int(9.3 * 914400)

    # ---- judges' guide: numbers after the three new slides
    where = {"TextBox 554": "الشريحتان 15 و16", "TextBox 559": "الشريحة 21", "TextBox 564": "الشرائح 5 إلى 7",
             "TextBox 569": "الشرائح 8 إلى 11", "TextBox 574": "الشرائح 4 و12 و18 إلى 20",
             "TextBox 579": "الشريحتان 23 و25 والملحق", "TextBox 584": "الشرائح 13 و14 و17"}
    for name, text in where.items():
        fill(guide, by_name(guide, name), [text])

    # ---- table of contents
    toc = base(prs, guide)
    kicker(toc, "فهرس العرض")
    fill(toc, ph(toc, 0), ["المحتويات"])
    remove(ph(toc, 2))
    items = [("01", "المشكلة", "4"), ("02", "الحل والمقارنة", "5 إلى 7"), ("03", "تجربة الطفل", "8 و9"),
             ("04", "تجربة الوالدين والميزات", "10 و11"), ("05", "القيم", "12"), ("06", "مصادر موثوقة", "13 و14"),
             ("07", "كيف يعمل والتقنية", "15 و16"), ("08", "السلامة", "17"),
             ("09", "النتائج والأثر المتوقع", "18 و19"), ("10", "ما بنيناه", "20"), ("11", "التشغيل", "21"),
             ("12", "الحدود والخطوة التالية", "22"), ("13", "الحسابات وتجربة اللجنة", "23"), ("14", "الفريق", "24"),
             ("15", "تحقق بنفسك", "25"), ("+", "الملحق: الأدلة التفصيلية", "27")]
    cw, rh = (W0 - 0.6) / 2, 0.74
    for i, (n, name, sl) in enumerate(items):
        col, row = divmod(i, 8)
        x = X0 + W0 - (col + 1) * cw - col * 0.6
        y = 3.3 + row * rh
        add_box(toc, x, y + 0.06, cw, rh - 0.12, fill=CARD, line=CARDLINE, radius=0.25)
        add_box(toc, x + cw - 0.82, y + 0.12, 0.62, rh - 0.24, fill=TURQ, radius=0.5)
        add_text(toc, x + cw - 0.82, y + 0.12, 0.62, rh - 0.24, [n], size=MIN_PT, bold=True, color=NAVY,
                 align="c", anchor="m")
        add_text(toc, x + 2.0, y, cw - 3.0, rh, [name], size=22, bold=True, color=OFF, anchor="m")
        add_text(toc, x + 0.25, y, 1.7, rh, [f"شريحة {sl}" if sl.isdigit() else f"الشرائح {sl}"], size=MIN_PT,
                 color=TURQ, align="l", anchor="m")
    notes(toc, "فهرس العرض: خمسة عشر قسما ثم الملحق. كل قسم يحمل في أعلى شرائحه رقمه واسم معيار التحكيم الذي يجيب عنه، "
               "وشريحة دليل اللجنة التالية تربط كل معيار بشرائحه.")

    # ---- comparison with current AI companions
    cmp_ = base(prs, guide)
    kicker(cmp_, "02 · الحل · معيار الابتكار")
    fill(cmp_, ph(cmp_, 0), ["مقارنة بالرفقاء الحاليين"])
    fill(cmp_, ph(cmp_, 2), ["ChatGPT روبوت محادثة لا رفيق، وGrok وReplika رفقاء للكبار"])
    heads = ["الصديق الصدوق", "ChatGPT", "Grok Companions", "Replika"]
    Y, N, P = "FaCircleCheck", "FaCircleXmark", "FaCircleHalfStroke"
    rows = [("العمر المسموح", ["6 إلى 13", "13+", "13+", "18+"]),
            ("رفيق بشخصية وصوت", [Y, N, Y, Y]),
            ("جواب ديني بمصدر، ولا فتوى", [Y, N, N, N]),
            ("محتوى آمن للطفل بالتصميم", [Y, P, N, N]),
            ("ملخصات وتنبيهات للوالدين", [Y, P, N, N]),
            ("تعلّم القيم باللعب", [Y, N, N, N])]
    lw = 5.0
    cw = (W0 - lw) / 4
    top, rh = 3.9, 0.76
    xs = [X0 + W0 - lw - (k + 1) * cw for k in range(4)]
    add_box(cmp_, xs[0] + 0.06, top - 0.1, cw - 0.12, rh * (len(rows) + 1) + 0.15, fill=TURQ, fill_alpha=12,
            line=TURQ, radius=0.06)
    for k, h in enumerate(heads):
        add_text(cmp_, xs[k], top, cw, rh, [h], size=20, bold=True, color=TURQ if k == 0 else OFF, align="c",
                 anchor="m", rtl=k == 0)
    for r, (crit, vals) in enumerate(rows):
        y = top + (r + 1) * rh
        add_box(cmp_, X0, y + 0.04, W0, 0.02, fill=CARDLINE, shape=MSO_SHAPE.RECTANGLE)
        add_text(cmp_, X0 + W0 - lw, y, lw - 0.2, rh, [crit], size=20, color=OFF, anchor="m")
        for k, v in enumerate(vals):
            if v.startswith("Fa"):
                color = "turq" if v == Y else ("white" if v == P else "muted")
                d = 0.5
                cmp_.shapes.add_picture(str(__import__("restructure").ICONS / f"{v}-{color}.png"),
                                        int((xs[k] + cw / 2 - d / 2) * 914400), int((y + rh / 2 - d / 2) * 914400),
                                        int(d * 914400), int(d * 914400))
            else:
                add_text(cmp_, xs[k], y, cw, rh, [v], size=20, bold=k == 0, color=TURQ if k == 0 else BODY,
                         align="c", anchor="m", rtl=k == 0)
    add_text(cmp_, X0, 9.45, W0, 0.5, ["◐ جزئي · حسب الشروط والميزات المعلنة حتى أكتوبر 2026"], size=MIN_PT, color=MUTED)
    notes(cmp_, "قراءة الفريق للشروط والميزات المعلنة حتى أكتوبر 2026. "
                "ChatGPT: روبوت محادثة عام لا رفيق، الحد الأدنى للعمر 13 بموافقة الوالدين لمن دون 18، وفيه أدوات رقابة أبوية للمراهقين "
                "(ربط الحسابات، تجربة مناسبة للعمر، تنبيه عند ضيق شديد)، فوضعنا له جزئيا. "
                "Grok Companions: شخصيات متحركة من xAI منذ يوليو 2025، وشروط Grok تحدد 13 عاما دون تحقق من العمر، "
                "ورفيقة Ani فيها وضع للكبار، وأفادت تقارير بأن «وضع الأطفال» لا يمنعها. "
                "Replika: رفيق للكبار، 18 عاما فأكثر. "
                "لا يقدم أي منها جوابا دينيا من بنك مراجع، ولا تعلّم القيم بالمهام والأوسمة، ولا يُطلع الوالدين على مصادر ما سمعه الطفل.")

    # ---- expected impact (not measured yet)
    imp = base(prs, guide)
    kicker(imp, "09 · الأثر المتوقع · معيار تحقيق النفع")
    fill(imp, ph(imp, 0), ["الأثر الذي نتوقعه"])
    fill(imp, ph(imp, 2), ["أثر متوقع نقيسه في تجربة مع الأسر، لا نتيجة مقيسة بعد"])
    rings = [("FaChild", "الطفل", ["يسأل بأمان عن كل شيء", "يمارس قيمة كل يوم", "يعرف مصدر ما يتعلمه"]),
             ("FaPeopleRoof", "الأسرة", ["حوار أكثر مع الوالدين", "والدان مطمئنان بلا تجسس", "مهام يشارك فيها الوالدان"]),
             ("FaUsers", "المجتمع", ["جيل يتحقق من المصدر", "محتوى إسلامي موثوق للأطفال", "بديل آمن عن الرفقاء العامين"])]
    tw, gap = 4.9, 1.09
    for i, (ic, head, lines) in enumerate(rings):
        x = X0 + W0 - (i + 1) * tw - i * gap
        add_box(imp, x, 4.0, tw, 4.15, fill=CARD, line=TURQ if i == 1 else CARDLINE, radius=0.07)
        icon(imp, ic, x + tw / 2 - 0.6, 4.25, 1.2)
        add_text(imp, x + 0.2, 5.55, tw - 0.4, 0.65, [head], size=26, bold=True, color=TURQ, align="c")
        add_text(imp, x + 0.3, 6.3, tw - 0.6, 1.8, lines, size=MIN_PT, color=OFF, align="c")
        if i:
            add_arrow(imp, x + tw + 0.18, 5.85, gap - 0.36, 0.45)
    add_box(imp, X0, 8.5, W0, 1.0, fill=PURPLE, fill_alpha=25, line=PURPLE, radius=0.2)
    add_text(imp, X0 + 0.4, 8.5, W0 - 0.8, 1.0,
             [[{"text": "كيف سنقيسه: ", "bold": True, "color": TURQ},
               {"text": "القيم التي يمارسها الطفل، والمهام التي يؤكدها الوالدان، والحوارات التي يبدؤونها"}]],
             size=20, color=OFF, align="c", anchor="m")
    notes(imp, "هذه توقعات لا نتائج: لم نجرّب النظام بعد مع أسر حقيقية. على الطفل: مكان آمن للسؤال عن كل شيء، وممارسة يومية للقيم بالمهام، "
               "ومعرفة مصدر كل آية وحديث. على الأسرة: أسئلة للنقاش من ملخص كل جلسة تفتح حوارا بين الطفل ووالديه، وتنبيهات حين يلزم "
               "دون كشف نص المحادثة، ومهام يؤكدها الوالدان. على المجتمع: جيل يسأل عن مصدر المعلومة، ومحتوى إسلامي مراجع بلغة الطفل، "
               "وبديل آمن عن رفقاء الذكاء الاصطناعي العامين. نقيس ذلك في تجربة مع أسر بما يسجله النظام أصلا: القيم الممارسة، والمهام المؤكدة، "
               "واستعمال أسئلة النقاش.")

    # ---- safety: say what was tested, and what the two numbers mean
    fill(safety, ph(safety, 0), ["السلامة: كيف يحمي الصديق الطفل؟"])
    for sh in safety.shapes:
        if not sh.name.startswith("Google Shape;"):
            sh.top += int(0.32 * 914400)
    add_text(safety, X0, 3.08, W0, 0.45,
             ["قواعد ثابتة في كل محادثة، اختبرناها بـ42 عبارة بلسان أطفال، كل منها 3 مرات"], size=MIN_PT, color=BODY)
    for big, lab, name_big, name_lab in [("60/60", ["مشكلة حقيقية اكتشفها", "مثل تنمّر أو خوف من شخص"], "TextBox 567", "TextBox 568"),
                                        ("3/66", ["تنبيه خاطئ فقط للوالدين", "عند مشاعر عادية كالإحراج"], "TextBox 570", "TextBox 571")]:
        fill(safety, by_name(safety, name_big), [big])
        lb = by_name(safety, name_lab)
        lb.height = int(0.85 * 914400)
        fill(safety, lb, [{"text": lab[0], "size": 20, "bold": True}, {"text": lab[1], "size": MIN_PT}])

    # ---- verify: the demo video link replaces the team slot
    verify = S[21]
    fill(verify, ph(verify, 5), ["فيديو العرض"])
    for name in ("Rounded Rectangle 550", "TextBox 551"):
        remove(by_name(verify, name))
    add_text(verify, 1.88, 5.09, 7.60, 0.6,
             [{"text": "youtu.be/YYWhZCzcJgg", "link": "https://youtu.be/YYWhZCzcJgg", "bold": True}],
             size=22, color=TURQ, rtl=False)
    add_text(verify, 1.88, 5.7, 7.60, 0.5, ["عرض المنتج على يوتيوب"], size=MIN_PT, color=BODY)

    # ---- judge accounts file (Google Drive) on the verify and the accounts slides
    drive = "https://drive.google.com/file/d/170psUxzianwj8WSywF9Lc4RTxSRJ1YCC/view?usp=sharing"
    add_text(verify, 1.88, 6.35, 7.60, 0.6,
             [[{"text": "حسابات اللجنة: ", "bold": True, "color": OFF},
               {"text": "ملف الدخول على Google Drive", "link": drive, "bold": True}]],
             size=20, color=TURQ)
    accounts = S[19]
    add_text(accounts, X0, 9.25, 7.2, 0.5,
             [[{"text": "بيانات الدخول: ", "color": OFF}, {"text": "ملف حسابات اللجنة", "link": drive, "bold": True}]],
             size=MIN_PT, color=TURQ)

    # ---- order: contents after the cover, comparison after the solution, impact after the results
    lst = prs.slides._sldIdLst
    els = list(lst)[:len(S)]
    by_id = {int(el.get("id")): el for el in lst}
    new = {k: by_id[s.slide_id] for k, s in (("toc", toc), ("cmp", cmp_), ("imp", imp))}
    order = [els[0], new["toc"]] + els[1:4] + [new["cmp"]] + els[4:16] + [new["imp"]] + els[16:]
    for el in list(lst):
        lst.remove(el)
    for el in order:
        lst.append(el)
    # ---- keep the file under 20 MB: downscale oversized photos (the lead's 4069 px team photo)
    import io
    from PIL import Image
    for part in prs.part.package.iter_parts():
        if part.content_type in ("image/jpeg", "image/png") and len(part.blob) > 2_000_000:
            im = Image.open(io.BytesIO(part.blob))
            if max(im.size) > 1600:
                im.thumbnail((1600, 1600))
                buf = io.BytesIO()
                if part.content_type == "image/jpeg":
                    im.convert("RGB").save(buf, "JPEG", quality=85, optimize=True)
                else:
                    im.save(buf, "PNG", optimize=True)
                part._blob = buf.getvalue()
    prs.save(OUTFILE)
    print("wrote", OUTFILE, len(order), "slides")


if __name__ == "__main__":
    build(sys.argv[1])

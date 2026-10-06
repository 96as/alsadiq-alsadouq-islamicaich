"""Dump a pptx: slide size, layouts, and every slide's shapes with text/geometry."""
import sys
from pptx import Presentation
from pptx.util import Emu

path = sys.argv[1]
prs = Presentation(path)
E = 914400
print("size", prs.slide_width / E, prs.slide_height / E, "in")
for mi, m in enumerate(prs.slide_masters):
    print(f"MASTER {mi}: {len(m.slide_layouts)} layouts")
    for li, l in enumerate(m.slide_layouts):
        print(f"  L{li}: {l.name!r} shapes={len(l.shapes)} ph={[ (p.placeholder_format.idx, p.placeholder_format.type, p.name) for p in l.placeholders]}")

only = None
if len(sys.argv) > 2:
    only = set(int(x) for x in sys.argv[2].split(","))

for si, s in enumerate(prs.slides, 1):
    if only and si not in only:
        continue
    print(f"\n=== SLIDE {si} layout={s.slide_layout.name!r}")
    def walk(shapes, depth=0):
        for sh in shapes:
            pad = "  " * (depth + 1)
            geo = ""
            try:
                geo = f"@({sh.left/E:.2f},{sh.top/E:.2f}) {sh.width/E:.2f}x{sh.height/E:.2f}"
            except Exception:
                pass
            kind = sh.shape_type
            extra = ""
            if sh.is_placeholder:
                extra = f" PH idx={sh.placeholder_format.idx}"
            print(f"{pad}- {sh.name} [{kind}]{extra} {geo}")
            if sh.has_text_frame if hasattr(sh, 'has_text_frame') else False:
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs)
                    if t.strip():
                        sizes = {r.font.size.pt if r.font.size else None for r in p.runs}
                        fonts = {r.font.name for r in p.runs}
                        print(f"{pad}    txt: {t[:140]!r} sz={sizes} f={fonts} algn={p.alignment}")
            if kind == 6:  # group
                walk(sh.shapes, depth + 1)
            if getattr(sh, 'has_table', False) and sh.has_table:
                for r in sh.table.rows:
                    print(f"{pad}    row: {[c.text[:30] for c in r.cells]}")
    walk(s.shapes)

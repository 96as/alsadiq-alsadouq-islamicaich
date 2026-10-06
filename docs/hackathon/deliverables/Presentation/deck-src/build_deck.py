"""Build the final Al-Sadiq Al-Sadouq deck on the organisers' template.

    python build_deck.py            -> Al-Sadiq-Al-Sadouq-Final.pptx

Method: every slide is a clone of one of the template's sample slides (so the
background art, logos, footer and placeholder styles stay the organisers'),
filled with our text, then the guide slides, the original samples and the
unused layouts are removed. Numbers carry their source and date on the slide.
No Quran or hadith text is typed here; the only scripture on any slide is
inside screenshots of the app's own source cards.

Structure: a main deck of 14 slides for the 5-minute pitch, in the template's
order (problem, solution, how it works, prototype, impact, team), each main
evidence slide naming the judging criterion it answers; then an appendix
divider and the detailed evidence slides for the panel's questions.

Inputs (not in git): organiser-template.pptx and img/ (PNG copies of the
README screenshots, made by prep_images.py); compose.py reads the parent
sources screenshot from the repository's readme-media folder.
"""
import copy
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

import compose
import cost_estimate as ce

HERE = Path(__file__).parent
TEMPLATE = HERE / "organiser-template.pptx"
OUTFILE = HERE / "Al-Sadiq-Al-Sadouq-Final.pptx"
E = 914400

NAVY, PURPLE, TURQ, OFF = "12183F", "6150EA", "2EF2C2", "F2F4FF"
BODY, MUTED, CARD, CARDLINE, LIGHTMUTED = "CBCFE6", "8E99CC", "161D48", "4A5590", "4A5590"
PILL = "143C23"   # the dark green of the state label burnt into the avatar frames
FONT = "Readex Pro"
DEMO_URL = "https://alsadiqai.com"
MIN_PT = 18       # no text of ours is smaller than the template's caption size


def I(x):
    return int(round(x * E))


# ---------------------------------------------------------------- the USD estimate (one source: cost_estimate.py)

def _cost_strings():
    # every figure is rounded UP at its display step (ce.display), so no shown bound is below the computed value
    shown = ce.display()
    lo_items, _ = ce.session("low")
    (j_lo, j_hi), (f_lo, f_hi) = ce.SCENARIOS.values()
    rng = lambda k: f"{shown[k][0]} إلى {shown[k][1]}"
    return {
        "items": {k: shown[k] for k in lo_items},
        "session": rng("session"),
        "judges": rng("judges"),
        "families": rng("families"),
        "judge_sessions": (j_lo, j_hi),
        "family_sessions": (f_lo, f_hi),
    }


EST = _cost_strings()


# ---------------------------------------------------------------- slide plumbing

def clone_slide(prs, src):
    """Copy a template slide (shapes, pictures, background) to a new slide at the end."""
    new = prs.slides.add_slide(src.slide_layout)
    tree = new.shapes._spTree
    for el in list(tree):
        if el.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            tree.remove(el)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rel.is_external:
            rid_map[rid] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rid] = new.part.relate_to(rel.target_part, rel.reltype)
    for el in src.shapes._spTree:
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        c = copy.deepcopy(el)
        for node in c.iter():
            for attr in (qn("r:embed"), qn("r:id"), qn("r:link")):
                v = node.get(attr)
                if v and v in rid_map:
                    node.set(attr, rid_map[v])
        tree.append(c)
    bg = src._element.cSld.find(qn("p:bg"))
    if bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(bg))
    return new


def ph(slide, idx, nth=0):
    found = [s for s in slide.shapes if s.is_placeholder and s.placeholder_format.idx == idx]
    return found[nth] if len(found) > nth else None


def by_name(slide, needle):
    for s in slide.shapes:
        if s.name.endswith(needle):
            return s
    raise KeyError(needle)


def remove(shape):
    if shape is not None:
        shape._element.getparent().remove(shape._element)


def geom(shape):
    return shape.left / E, shape.top / E, shape.width / E, shape.height / E


def set_geom(shape, x=None, y=None, w=None, h=None):
    if x is not None:
        shape.left = I(x)
    if y is not None:
        shape.top = I(y)
    if w is not None:
        shape.width = I(w)
    if h is not None:
        shape.height = I(h)


# ---------------------------------------------------------------- text

def _rpr(proto, size=None, bold=None, color=None, light=False):
    rpr = copy.deepcopy(proto) if proto is not None else etree.Element(qn("a:rPr"))
    rpr.set("lang", "ar-SA")
    if size is not None:
        rpr.set("sz", str(int(round(size * 100))))
    if bold is not None:
        rpr.set("b", "1" if bold else "0")
    if color is not None:
        for f in rpr.findall(qn("a:solidFill")):
            rpr.remove(f)
        sf = etree.Element(qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
        # solidFill must precede latin/ea/cs
        rpr.insert(0, sf)
    face = FONT + " Light" if light else None
    for tag in ("a:latin", "a:ea", "a:cs", "a:sym"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", face or el.get("typeface") or FONT)
        if face is None and el.get("typeface") not in (FONT, FONT + " Light", FONT + " Medium", FONT + " SemiBold"):
            el.set("typeface", FONT)
    return rpr


def _add_link(slide, rpr, url):
    rid = slide.part.relate_to(url, RT.HYPERLINK, is_external=True)
    for old in rpr.findall(qn("a:hlinkClick")):
        rpr.remove(old)
    h = etree.SubElement(rpr, qn("a:hlinkClick"))
    h.set(qn("r:id"), rid)


def fill(slide, shape, paras, **default):
    """Replace a shape's paragraphs, keeping the template's pPr/rPr per paragraph.

    paras: list of paragraphs; a paragraph is a str, a dict(text=..., size=, bold=, color=, link=, light=),
    or a list of such run dicts (several runs in one paragraph).
    """
    tx = shape.text_frame._txBody
    old = tx.findall(qn("a:p"))
    protos = []
    for p in old:
        ppr = p.find(qn("a:pPr"))
        r = p.find(qn("a:r"))
        if r is None:
            r = p.find(qn("a:fld"))
        rpr = r.find(qn("a:rPr")) if r is not None else None
        protos.append((copy.deepcopy(ppr), copy.deepcopy(rpr)))
    if not protos:
        protos = [(None, None)]
    for p in old:
        tx.remove(p)
    for i, para in enumerate(paras):
        ppr, rpr = protos[min(i, len(protos) - 1)]
        p = etree.SubElement(tx, qn("a:p"))
        if ppr is not None:
            p.append(copy.deepcopy(ppr))
        runs = para if isinstance(para, list) else [para]
        for run in runs:
            if isinstance(run, str):
                run = {"text": run}
            opts = dict(default)
            opts.update(run)
            r = etree.SubElement(p, qn("a:r"))
            rp = _rpr(rpr, opts.get("size"), opts.get("bold"), opts.get("color"), opts.get("light", False))
            if opts.get("link"):
                _add_link(slide, rp, opts["link"])
            r.append(rp)
            etree.SubElement(r, qn("a:t")).text = opts["text"]
    return shape


def add_text(slide, x, y, w, h, paras, size=24, color=OFF, bold=False, align="r", anchor="t", rtl=True,
             light=False, spacing=None):
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    body = tf._txBody.find(qn("a:bodyPr"))
    for k in ("lIns", "rIns", "tIns", "bIns"):
        body.set(k, "0")
    body.set("anchor", {"t": "t", "m": "ctr", "b": "b"}[anchor])
    body.set("wrap", "square")
    for p in tf._txBody.findall(qn("a:p")):
        tf._txBody.remove(p)
    for para in paras:
        runs = para if isinstance(para, list) else [para]
        p = etree.SubElement(tf._txBody, qn("a:p"))
        ppr = etree.SubElement(p, qn("a:pPr"))
        ppr.set("algn", {"r": "r", "l": "l", "c": "ctr"}[align])
        ppr.set("rtl", "1" if rtl else "0")
        if spacing:
            sb = etree.SubElement(ppr, qn("a:spcBef"))
            etree.SubElement(sb, qn("a:spcPts")).set("val", str(int(spacing * 100)))
        for run in runs:
            if isinstance(run, str):
                run = {"text": run}
            r = etree.SubElement(p, qn("a:r"))
            rp = _rpr(None, run.get("size", size), run.get("bold", bold), run.get("color", color),
                      run.get("light", light))
            if run.get("link"):
                _add_link(slide, rp, run["link"])
            r.append(rp)
            etree.SubElement(r, qn("a:t")).text = run["text"]
    return tb


def _fill_xml(spPr, color, alpha=None):
    for f in spPr.findall(qn("a:solidFill")) + spPr.findall(qn("a:noFill")):
        spPr.remove(f)
    sf = etree.Element(qn("a:solidFill"))
    c = etree.SubElement(sf, qn("a:srgbClr"))
    c.set("val", color)
    if alpha is not None:
        etree.SubElement(c, qn("a:alpha")).set("val", str(int(alpha * 1000)))
    geom_el = spPr.find(qn("a:prstGeom"))
    geom_el.addnext(sf)


def _line_xml(spPr, color, width_pt=1.0, alpha=None, dash=None):
    for ln in spPr.findall(qn("a:ln")):
        spPr.remove(ln)
    ln = etree.SubElement(spPr, qn("a:ln"))
    ln.set("w", str(int(width_pt * 12700)))
    if color is None:
        etree.SubElement(ln, qn("a:noFill"))
        return
    sf = etree.SubElement(ln, qn("a:solidFill"))
    c = etree.SubElement(sf, qn("a:srgbClr"))
    c.set("val", color)
    if alpha is not None:
        etree.SubElement(c, qn("a:alpha")).set("val", str(int(alpha * 1000)))
    if dash:
        etree.SubElement(ln, qn("a:prstDash")).set("val", dash)


def add_box(slide, x, y, w, h, fill=None, fill_alpha=None, line=None, line_alpha=None, line_w=1.0, dash=None,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    spPr = s._element.spPr
    if fill:
        _fill_xml(spPr, fill, fill_alpha)
    else:
        for f in spPr.findall(qn("a:solidFill")):
            spPr.remove(f)
        spPr.find(qn("a:prstGeom")).addnext(etree.Element(qn("a:noFill")))
    _line_xml(spPr, line, line_w, line_alpha, dash)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    # no shadow/style text
    st = s._element.find(qn("p:style"))
    if st is not None:
        s._element.remove(st)
    s.text_frame.text = ""
    return s


def add_slot(slide, x, y, w, h, label, dark=True, size=MIN_PT, solid=False):
    """A clearly marked slot for a teammate to fill (dashed box + label)."""
    accent = TURQ if dark else PURPLE
    if solid:
        add_box(slide, x, y, w, h, fill=NAVY, fill_alpha=92, radius=0.06)
    add_box(slide, x, y, w, h, fill=accent, fill_alpha=7, line=accent, line_w=2.0, dash="dash", radius=0.06)
    add_text(slide, x + 0.15, y + 0.08, w - 0.3, h - 0.16,
             [[{"text": "خانة للفريق: ", "bold": True, "color": accent},
               {"text": label, "color": OFF if dark else NAVY}]],
             size=size, anchor="m", align="c")


def add_picture(slide, path, x, y, w, h, radius=None):
    pic = slide.shapes.add_picture(str(path), I(x), I(y), I(w), I(h))
    if radius:
        g = pic._element.spPr.find(qn("a:prstGeom"))
        g.set("prst", "roundRect")
        av = g.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(g, qn("a:avLst"))
        gd = etree.SubElement(av, qn("a:gd"))
        gd.set("name", "adj")
        gd.set("fmla", f"val {int(radius * 100000)}")
    return pic


def add_arrow(slide, x, y, w, h, kind=MSO_SHAPE.LEFT_ARROW, color=TURQ, alpha=80):
    s = slide.shapes.add_shape(kind, I(x), I(y), I(w), I(h))
    _fill_xml(s._element.spPr, color, alpha)
    _line_xml(s._element.spPr, None)
    st = s._element.find(qn("p:style"))
    if st is not None:
        s._element.remove(st)
    return s


def top_anchor(shape, top_in=0.55):
    body = shape.text_frame._txBody.find(qn("a:bodyPr"))
    body.set("anchor", "t")
    body.set("tIns", str(I(top_in)))
    return shape


def set_theme_links(prs, color=TURQ):
    for master in prs.slide_masters:
        theme = master.part.part_related_by(RT.THEME)
        xml = theme.blob.decode("utf8")
        xml = re.sub(r"(<a:hlink>\s*<a:srgbClr val=\")[0-9A-Fa-f]{6}", r"\g<1>" + color, xml)
        xml = re.sub(r"(<a:folHlink>\s*<a:srgbClr val=\")[0-9A-Fa-f]{6}", r"\g<1>" + color, xml)
        theme._blob = xml.encode("utf8")


def caption(slide, text, y=9.62, dark=True, x=1.54, w=16.88, size=MIN_PT, h=0.6):
    return add_text(slide, x, y, w, h, [text], size=size, color=MUTED if dark else LIGHTMUTED, light=False)


def kicker(slide, text):
    fill(slide, ph(slide, 1), [text])


# ---------------------------------------------------------------- tables

def fill_table(slide, gf, rows, col_widths_rtl, row_h=0.62, size=20, header_size=22.5, row_heights=None):
    """rows are written right to left (as read); the template table is LTR internally.

    row_heights, when given, is one height per row and overrides row_h."""
    tbl = gf.table._tbl
    grid = tbl.find(qn("a:tblGrid"))
    cols = list(reversed(col_widths_rtl))
    for gc, w in zip(grid.findall(qn("a:gridCol")), cols):
        gc.set("w", str(I(w)))
    trs = tbl.findall(qn("a:tr"))
    header_proto, body_proto = trs[0], trs[1]
    for tr in trs:
        tbl.remove(tr)
    for ri, row in enumerate(rows):
        tr = copy.deepcopy(header_proto if ri == 0 else body_proto)
        tr.set("h", str(I(row_heights[ri] if row_heights else row_h)))
        tbl.append(tr)
        cells = tr.findall(qn("a:tc"))
        for ci, text in enumerate(reversed(row)):
            tc = cells[ci]
            txb = tc.find(qn("a:txBody"))
            p = txb.find(qn("a:p"))
            r = p.find(qn("a:r"))
            rpr = copy.deepcopy(r.find(qn("a:rPr")))
            for rr in p.findall(qn("a:r")):
                p.remove(rr)
            end = p.find(qn("a:endParaRPr"))
            spec = text if isinstance(text, dict) else {"text": text}
            sz = spec.get("size", header_size if ri == 0 else size)
            rp = _rpr(rpr, sz, spec.get("bold"), spec.get("color"))
            # "\n" in a cell's text is a line break inside the cell (a:br between runs)
            for li, part in enumerate(spec["text"].split("\n")):
                if li:
                    br = etree.Element(qn("a:br"))
                    br.append(copy.deepcopy(rp))
                    if end is not None:
                        end.addprevious(br)
                    else:
                        p.append(br)
                nr = etree.Element(qn("a:r"))
                nr.append(copy.deepcopy(rp))
                etree.SubElement(nr, qn("a:t")).text = part
                if end is not None:
                    end.addprevious(nr)
                else:
                    p.append(nr)
    total_w = sum(col_widths_rtl)
    gf.width = I(total_w)
    gf.height = I(sum(row_heights) if row_heights else row_h * len(rows))


# ---------------------------------------------------------------- build

def build():
    imgs = compose.build()
    prs = Presentation(str(TEMPLATE))
    originals = list(prs.slides)
    orig_ids = list(prs.slides._sldIdLst)
    S = {i + 1: s for i, s in enumerate(originals)}   # template slide number -> slide

    def new(n):
        return clone_slide(prs, S[n])

    # ================================================================ MAIN DECK (5-minute pitch)

    # M1 Cover -----------------------------------------------------------
    s = new(8)
    fill(s, ph(s, 0), ["الصديق الصدوق"])
    fill(s, ph(s, 1), ["صديق صوتي يعلم الأطفال 38 قيمة إسلامية، ويريهم المصدر الموثق حين يسألون"])
    team = by_name(s, ";489;p8")
    set_geom(team, x=12.4, w=6.01)
    fill(s, team, ["فريق الصديق الصدوق · المسار المفتوح"])
    fill(s, by_name(s, ";490;p8"), ["6 أكتوبر 2026"])

    # M2 Problem (four points, dark) ----------------------------------------
    s = new(11)
    kicker(s, "01 · المشكلة")
    fill(s, ph(s, 0), ["الأطفال يسألون الذكاء الاصطناعي عن دينهم"])
    fill(s, ph(s, 2), ["والمساعدات العامة لم تصمم لطفل، ولا تعرف مصدر ما تقوله"])
    pts = [(3, 4, "إجابة بلا مصدر", "قد يخترع حديثا أو يخطئ في آية، ولا يذكر من أين جاء بها"),
           (5, 6, "فتوى لطفل", "يجيب عن مسائل شخصية في الفقه تحتاج والدا أو عالما"),
           (7, 8, "والدان خارج الحوار", "لا يعرفان ما قيل لطفلهما، ولا يستطيعان التحقق منه"),
           (9, 13, "محتوى للكبار", "المراجع الموثوقة مكتوبة للكبار، لا بلغة طفل في السادسة")]
    for a, b, t, d in pts:
        fill(s, ph(s, a), [t])
        fill(s, ph(s, b), [d])

    # M3 Solution (text + wide image) ---------------------------------------
    s = new(13)
    kicker(s, "02 · الحل")
    fill(s, ph(s, 0), ["الصديق الصدوق: صديق يحكي، ومصدر حين يسأل"])
    p = ph(s, 2)
    x, y, w, h = geom(p)
    remove(p)
    sw = h * imgs["solution_aspect"]          # the home screen at its own aspect, demo bar removed
    add_picture(s, imgs["solution"], x, y, sw, h, radius=0.04)
    fill(s, ph(s, 3), ["ما هو؟"])
    fill(s, ph(s, 4), ["صديق صوتي بالعربية والإنجليزية للأطفال من 6 إلى 13 سنة، يعلم 38 قيمة إسلامية بالحوار والمهام"])
    fill(s, ph(s, 5), ["ما الذي يميزه؟"])
    fill(s, ph(s, 6), ["لا يذكر آية أو حديثا إلا من بنك راجعه الفريق، ويريه للطفل ولوالديه بمصدره"])
    caption(s, "الشاشة الرئيسية للطفل بالعربية، لعائلة تجريبية ببيانات اصطناعية", y=9.72, x=x, w=sw)

    # M4 Innovation (identity icons) ------------------------------------------
    s = new(27)
    kicker(s, "02 · الحل · معيار الابتكار")
    fill(s, ph(s, 0), ["ما الجديد في الصديق الصدوق؟"])
    fill(s, ph(s, 2), ["ما يضيفه مقارنة بمساعد ذكاء اصطناعي عام أو تطبيق محتوى تقليدي"])
    inn = [(3, 4, "صديق لا واعظ", "يدردش ويفرح مع الطفل، ولا يقحم درسا لم يطلبه"),
           (5, 6, "مصدر حين يسأل", "بطاقة مراجعة تظهر حين يسأل الطفل عن دينه فقط"),
           (7, 8, "تلاوة لا صوت آلي", "الآيات بتلاوة مسجلة، ولا تمر على محول النص إلى صوت"),
           (9, 13, "المصدر للوالدين", "يريان كل مصدر ظهر لطفلهما هذا الأسبوع، برابطه")]
    for a, b, t, d in inn:
        fill(s, ph(s, a), [t])
        fill(s, ph(s, b), [d])

    # M5 Architecture (dark, custom diagram on the four-points base) -----------
    s = new(11)
    kicker(s, "03 · آلية العمل · معيار الجودة التقنية")
    fill(s, ph(s, 0), ["البنية التقنية للوكيل الصوتي"])
    fill(s, ph(s, 2), ["وكيل LiveKit يقرر قبل أن يتكلم: الأمان أولا، ثم الإحالة، ثم الاسترجاع من البنك"])
    for idx in (3, 4, 5, 6, 7, 8, 9, 13):
        remove(ph(s, idx))
    bw, gap, top, bh = 3.02, 0.445, 4.1, 2.0
    xs = [1.54 + (4 - i) * (bw + gap) for i in range(5)]   # i=0 is the rightmost box

    def node(x, y, w, title, ar, latin, hi=False):
        add_box(s, x, y, w, bh, fill=(TURQ if hi else "FFFFFF"), fill_alpha=(10 if hi else 5),
                line=TURQ, line_alpha=(100 if hi else 25), line_w=(2.0 if hi else 1.0))
        add_text(s, x + 0.12, y + 0.1, w - 0.24, bh - 0.2,
                 [{"text": title, "bold": True, "size": 21, "color": TURQ if hi else OFF},
                  {"text": ar, "size": MIN_PT, "color": BODY},
                  {"text": latin, "size": MIN_PT, "color": MUTED}], anchor="m", align="c", spacing=2)

    row1 = [("سؤال الطفل", "صوت أو كتابة", "LiveKit WebRTC"),
            ("الكلام إلى نص", "تفريغ صوت الطفل", "OpenAI"),
            ("سياسة الدور", "أمان، إحالة، استرجاع", "turn_policy.py"),
            ("صياغة الرد", "يقتبس من المصدر فقط", "gpt-5.4-mini"),
            ("صوت الرد", "لا يقرأ الآيات آليا", "ElevenLabs Flash")]
    for i, (t, ar, la) in enumerate(row1):
        node(xs[i], top, bw, t, ar, la, hi=(i == 2))
        if i < 4:
            add_arrow(s, xs[i] - gap + 0.06, top + bh / 2 - 0.16, gap - 0.12, 0.32)
    top2 = 7.0
    wide = 2 * bw + gap
    node(xs[1], top2, wide, "تطبيق الطفل", "شخصية بمزامنة الشفاه، وبطاقة المصدر، وتلاوة مسجلة",
         "React Three Fiber · everyayah")
    node(xs[2], top2, bw, "البنك المراجع", "273 عنصرا و38 قيمة", "PostgreSQL", hi=True)
    add_arrow(s, xs[2] + bw / 2 - 0.2, top + bh + 0.08, 0.4, top2 - top - bh - 0.16, kind=MSO_SHAPE.UP_DOWN_ARROW)
    add_text(s, xs[2] + bw / 2 + 0.3, top + bh + 0.2, 3.6, 0.5, ["استرجاع في نحو ميلي ثانية"],
             size=MIN_PT, color=MUTED)
    node(xs[4], top2, wide, "لوحة الوالدين", "سجل المصادر التي خدمها البنك فعلا، لا ما يقوله النموذج",
         "Django · sources_used")
    caption(s, "المصدر: الشيفرة المدمجة في فرع hackathon عند 735c95a · زمن الاسترجاع من build-progress.md في 5 أكتوبر 2026",
            y=9.3)

    # M6 Content bank table (dark) ---------------------------------------------
    s = new(17)
    kicker(s, "03 · آلية العمل · معيار الموثوقية")
    fill(s, ph(s, 0), ["بنك المحتوى: مصدر معتمد لكل عنصر"])
    fill(s, ph(s, 2), ["273 عنصرا كلها مراجعة ومرتبطة بـ38 قيمة، ولكل عنصر مصدره ورابطه"])
    gf = [sh for sh in s.shapes if sh.has_table][0]
    gf.left, gf.top = I(1.54), I(4.05)
    fill_table(s, gf, [
        ["النوع", "العدد", "المصدر المعتمد", "التحقق"],
        ["آيات", "74", "مجمع الملك فهد عبر Quranpedia، والتلاوة من everyayah", "نص منسوخ آليا ومثبت ببصمة"],
        ["أحاديث", "94", "الدرر السنية: البخاري 60 ومسلم 34، كلها صحيحة", "رابط الدرر، والإنجليزية لـ42 منها من HadeethEnc"],
        ["مصطلحات", "32", "معجم الجمهرة", "نسخ من الصفحة مع تحقق ببصمة"],
        ["أسئلة شائعة", "23", "كتاب بينات", "رقم السؤال والصفحة والرابط"],
        ["تفسير وعقيدة وفقه وسيرة", "50", "موسوعات الدرر السنية", "روجعت عنصرا عنصرا"],
    ], col_widths_rtl=[4.1, 1.5, 6.68, 4.6], row_h=0.78, size=20)
    src = ph(s, 3)
    set_geom(src, y=9.0, h=0.9)
    fill(s, src, ["المصدر: README القسم 2.3 وbank-counts.md · البنك عند 735c95a لم يتغير منذ 8c2ff0d، وseed_content يطبع: reviewed 273, rejected 0"])

    # M7 Prototype gallery ---------------------------------------------------
    s = new(26)
    kicker(s, "04 · النموذج الأولي · معيار تجربة المستفيد")
    fill(s, ph(s, 0), ["تجربة الطفل: يتكلم، فيسمع الصديق ويرد"])
    pics = [sh for sh in s.shapes if sh.is_placeholder and sh.placeholder_format.idx == 2]
    pics.sort(key=lambda sh: -sh.left)   # right to left
    for sh, key in zip(pics, ["g_home", "g_voice", "g_chat"]):
        x, y, w, h = geom(sh)
        remove(sh)
        add_picture(s, imgs[key], x, y, w, h, radius=0.05)
    fill(s, ph(s, 3), ["البداية بزر واحد"])
    fill(s, ph(s, 5), ["مكالمة صوتية مع الصديق"])
    fill(s, ph(s, 7), ["الكتابة بديلا، مع بطاقة المصدر"])
    caption(s, "عربي أولا مع الإنجليزية · الكتابة بديل كامل للصوت · 0 مخالفات في فحص axe على 12 صفحة باللغتين (product-web.md)",
            y=9.72, h=0.5)

    # M8 Source cards (large image) -------------------------------------------
    s = new(25)
    p = ph(s, 2)
    x, y, w, h = geom(p)
    remove(p)
    add_picture(s, imgs["cards"], x, y, w, h)
    fill(s, ph(s, 0), ["بطاقة المصدر: ما يراه الطفل حين يسأل"])
    fill(s, ph(s, 1), ["من التطبيق: آية بتلاوة مسجلة ورابط Quranpedia، وحديث من الدرر بدرجته، ومصطلح من الجمهرة، وسؤال من بينات"])

    # M9 Parent view: the merged "Sources discussed this week" page -----------
    s = new(13)
    kicker(s, "04 · النموذج الأولي")
    fill(s, ph(s, 0), ["الوالدان يريان المصدر نفسه"])
    p = ph(s, 2)
    x, y, w, h = geom(p)
    remove(p)
    add_picture(s, imgs["parent"], x, y, w, h)
    fill(s, ph(s, 3), ["ما الذي يراه الوالدان؟"])
    fill(s, ph(s, 4), ["كل مصدر ظهر لطفلهما هذا الأسبوع: نصه ومرجعه ودرجته ورابطه، وعدد مرات نقاشه"])
    fill(s, ph(s, 5), ["لماذا يثقان به؟"])
    fill(s, ph(s, 6), ["تسجل المصادر مما خدمه البنك فعلا لا مما قاله النموذج، ولا يظهر لهما نص المحادثة"])
    caption(s, "صفحة الرؤى لولي الأمر بالعربية، لعائلة تجريبية · PR #59 وPR #58", y=9.72, x=x, w=w, h=0.6)

    # M10 Results (KPIs) --------------------------------------------------
    s = new(15)
    kicker(s, "05 · الأثر · معيار تحقيق النفع")
    fill(s, ph(s, 0), ["نتائج مقيسة: قبل العمل وبعده"])
    fill(s, ph(s, 2), ["مقياس المسار المفتوح: النجاح الصارم على حالات التقييم نفسها، مع فحوص سياسة الدور دون نموذج"])
    # The 0% tile is the lead's policy proxy on his own probe sets (PR #56), not a judged rate and not
    # reproduced by us; the tile and the footnote say so. The judged gate, before and after, is on appendix slide 22.
    # Bidi: a date never follows a Latin token across a comma or dot ("735c95a، 6 أكتوبر" would pull
    # the day into the Latin run and print it on the wrong side); an Arabic word sits between them.
    kpis = [(3, 4, 5, "44%", "المجموعة المحجوزة", "نجاح صارم، وكان 25%|تقييم حي بمحكم آلي، 5 أكتوبر"),
            (6, 7, 8, "340", "اجتازت تقييم السياسة", "من 392 تشغيلا دون نموذج|عند 735c95a في 6 أكتوبر"),
            (9, 13, 14, "92/92", "تشغيلات أمان ناجحة", "في التقييم نفسه|ولم يفشل أي تشغيل للأمان"),
            (15, 16, 17, "0%", "مؤشر القائد التقريبي", "للسياسة على مجموعات اختباره|5 أكتوبر: كان 51.5% قبل PR #56")]
    for a, b, c, v, n, u in kpis:
        fill(s, ph(s, a), [v])
        fill(s, ph(s, b), [n])
        fill(s, ph(s, c), u.split("|"))
    # the proxy tile's label is longer: widen its name and unit boxes about the same centre
    for idx in (16, 17):
        box = ph(s, idx)
        bx, _, bw_, _ = geom(box)
        set_geom(box, x=bx - 0.25, w=bw_ + 0.5)
    src = ph(s, 18)
    set_geom(src, x=1.54, w=16.88, y=8.95, h=1.3)
    fill(s, src, ["المجموعة المحجوزة: 10 من 40 عند 4fccb39 ثم 18 من 41 عند dba765e، تشغيل واحد لكل جانب، p = 0.039 · eval-reports/README.md",
                  "سياسة الدور: run_eval دون نموذج · 0%: مؤشر القائد التقريبي لسياسة الدور على مجموعات اختباره، PR #56 في 5 أكتوبر",
                  "نسبة الأدوار العادية التي أظهرت بطاقة، منقولة عن وصف PR #56 في quality-gate-summary.md القسم 2 · لم نعد قياسها، وليست نسبة محكمة"])

    # M11 Operations and cost (dark table + estimate band) -----------------------
    # The table rows are measured; the USD band under it is an ESTIMATE (cost_estimate.py: measured units x
    # the providers' public prices of 6 Oct 2026). The band, the subtitle and the appendix say so.
    s = new(17)
    kicker(s, "05 · التشغيل · معيار واقعية التشغيل")
    fill(s, ph(s, 0), ["التشغيل والتكلفة"])
    fill(s, ph(s, 2), ["الحمل مقيس على نسخة محلية بخدمات OpenAI وElevenLabs الحقيقية في 5 أكتوبر 2026، والتكلفة بالدولار تقدير لا قياس"])
    gf = [sh for sh in s.shapes if sh.has_table][0]
    tbl_top, rh = 3.9, 0.5
    gf.left, gf.top = I(1.54), I(tbl_top)
    cols = [2.55, 7.18, 3.75, 3.4]
    fill_table(s, gf, [
        ["البند", "ما قيس", "القيمة", "المصدر"],
        ["السعة", "جلسات صوتية متزامنة على عامل وكيل واحد", "10 جلسات", "LOAD-REHEARSAL.md"],
        ["سرعة الرد", "من نهاية كلام الطفل إلى أول صوت للرد", "نحو 2.6 ثانية", "LOAD-REHEARSAL.md"],
        ["الموارد", "لكل جلسة عند الاستقرار", "0.1 نواة و165 ميغابايت", "LOAD-REHEARSAL.md"],
        ["الصوت", "أحرف ElevenLabs لجلسة قصيرة: تحية وسؤالان", "نحو 300 حرف", "LOAD-REHEARSAL.md"],
        ["النموذج", "استدعاءات OpenAI للجلسة نفسها، نصفها للتقرير", "نحو 8", "LOAD-REHEARSAL.md"],
        ["الرموز", "متوسط تشغيل التقييم (الوكيل والمحكم)، حد أعلى للاستدعاء", "إدخال 6.5 ألف، إخراج 109", "lead-summary.json"],
        ["النشر", "يطبق الترحيلات أولا، ويفشل إن لم يرد /api/health/", "آلي في كل نشر", "deploy.yml"],
    ], col_widths_rtl=cols, row_h=rh, size=19)
    # estimate band: a label on the right, then three figures, right to left
    band_y, band_h = tbl_top + rh * 8 + 0.15, 1.05
    lab_w, bgap = 3.3, 0.2
    add_box(s, 18.42 - lab_w, band_y, lab_w, band_h, fill=TURQ, fill_alpha=10, line=TURQ, line_w=1.5)
    add_text(s, 18.42 - lab_w + 0.15, band_y + 0.08, lab_w - 0.3, band_h - 0.16,
             [{"text": "التكلفة بالدولار", "bold": True, "size": 21, "color": TURQ},
              {"text": "تقدير، 6 أكتوبر 2026", "color": OFF},
              {"text": "المنهج في الملحق", "color": BODY}], size=MIN_PT, anchor="m")
    fig_w = (16.88 - lab_w - 3 * bgap) / 3
    figures = [(EST["session"], "دولار للجلسة الواحدة"),
               (EST["judges"], "دولارا لـ30 محكما طوال فترة التحكيم"),
               (EST["families"], "دولارا في الشهر لـ100 أسرة")]
    for i, (val, lab) in enumerate(figures):
        fx = 18.42 - lab_w - bgap - fig_w - i * (fig_w + bgap)
        add_box(s, fx, band_y, fig_w, band_h, fill="FFFFFF", fill_alpha=4, line=TURQ, line_alpha=35)
        add_text(s, fx + 0.15, band_y + 0.08, fig_w - 0.3, band_h - 0.16,
                 [{"text": val, "bold": True, "size": 25, "color": OFF},
                  {"text": lab, "color": BODY}], size=MIN_PT, anchor="m", align="c")
    src = ph(s, 3)
    set_geom(src, y=9.2, h=0.95)
    fill(s, src, ["الرموز من eval-reports/lead-summary.json عند f9ab61c في 5 أكتوبر · النشر: deploy.yml بعد PR #60 وPR #61 · الحماية من الإنفاق: سقف يومي لأحرف الصوت، وبعده تبدأ الجلسات الجديدة بالكتابة فقط (demo-guards.md)"])

    # M12 Team ---------------------------------------------------------------
    s = new(19)
    kicker(s, "06 · الفريق")
    fill(s, ph(s, 0), ["فريق العمل"])
    fill(s, ph(s, 2), ["فريق الصديق الصدوق: ثلاثة أعضاء"])
    for idx in (3, 4, 5, 6, 7, 8, 9, 13, 14, 15, 16, 17):
        remove(ph(s, idx))
    # the lead's answer of 6 Oct: the team slide lists these three members only
    members = [("ع", "عبدالرحمن سلامة", "قائد الفريق ومراجعة المحتوى", False),
               ("م", "مجد عوض", "الصوت والشخصية والواجهة والتقييم", False),
               ("ع", "عبدالرحمن محملجي", "مراجعة أقسام الدرر وصفحة المصادر", False)]
    colw, cgap = 4.9, 0.35
    right = 9.98 + (len(members) * colw + (len(members) - 1) * cgap) / 2   # the group centred on the slide
    for i, (ini, name, role, slot) in enumerate(members):
        cx = right - colw - i * (colw + cgap)
        add_box(s, cx + colw / 2 - 1.05, 4.2, 2.1, 2.1, fill=TURQ, fill_alpha=8, line=TURQ, line_w=2.0,
                shape=MSO_SHAPE.OVAL)
        add_text(s, cx + colw / 2 - 1.05, 4.2, 2.1, 2.1, [ini], size=54, bold=True, color=TURQ, align="c",
                 anchor="m")
        add_text(s, cx, 6.6, colw, 0.65, [name], size=23, bold=True, color=OFF, align="c")
        if slot:
            add_slot(s, cx + 0.1, 7.4, colw - 0.2, 1.3, role, size=MIN_PT)
        else:
            add_text(s, cx, 7.4, colw, 1.2, [role], size=20, color=TURQ, align="c")
    caption(s, "الأدوار من خطة المشروع PLAN.md ودفتر التقدم build-progress.md", y=9.55)

    # M13 Verify it yourself ----------------------------------------------
    s = new(11)
    kicker(s, "06 · التحقق · معيار وضوح العرض")
    fill(s, ph(s, 0), ["تحقق بنفسك"])
    fill(s, ph(s, 2), ["كل رقم في العرض معه مصدره وتاريخه، والشيفرة والتقارير في المستودع، والأدلة التفصيلية في الملحق"])
    fill(s, ph(s, 3), ["جرب الآن"])
    fill(s, ph(s, 4), [{"text": "alsadiqai.com", "link": DEMO_URL, "bold": True, "color": TURQ},
                       "ثم «جرب الصديق»: دون حساب، وببيانات تخيلية"])
    fill(s, ph(s, 5), ["الفيديو الاحتياطي"])
    b = ph(s, 6)
    x, y, w, h = geom(b)
    remove(b)
    add_slot(s, x, y + 0.05, w, 1.05, "رابط فيديو العرض، دقيقتان أو أقل")
    fill(s, ph(s, 7), ["المستودع والدليل README"])
    b = ph(s, 8)
    x, y, w, h = geom(b)
    remove(b)
    add_slot(s, x, y + 0.05, w, 1.05, "رابط المستودع العام بعد نشره")
    fill(s, ph(s, 9), ["الترخيص والإفصاح"])
    fill(s, ph(s, 13), ["جميع الحقوق محفوظة، والعمل السابق للتحدي معلن في DISCLOSURE.md"])

    # M14 Thank you -------------------------------------------------------
    s = new(31)
    fill(s, ph(s, 1), ["فريق الصديق الصدوق · المسار المفتوح"])
    fill(s, ph(s, 2), [{"text": "alsadiqai.com", "link": DEMO_URL}])

    # ================================================================ APPENDIX

    # Divider ---------------------------------------------------------------
    s = new(10)
    fill(s, ph(s, 1), ["ملحق"])
    fill(s, ph(s, 0), ["الأدلة التفصيلية"])
    fill(s, ph(s, 2), ["لأسئلة اللجنة: الأرقام بمصادرها ومنهجها"])
    add_text(s, 8.5, 8.3, 9.81, 1.2,
             ["آلية العمل · لا مصدر لا جواب · منهج المحتوى · سير المراجعة · الشخصية · الإتاحة",
              "بوابة الجودة · المؤشر التقريبي · تقدير التكلفة · الإفصاح · خارطة الطريق"],
             size=20, color=BODY, spacing=4)

    # A1 How it works -------------------------------------------------------
    s = new(24)
    kicker(s, "ملحق · آلية العمل")
    fill(s, ph(s, 0), ["كيف يعمل الصديق؟"])
    fill(s, ph(s, 2), ["من سؤال الطفل إلى مصدر يراه الوالدان، في كل دور صوتا أو كتابة"])
    steps = [(";822;p24", "01", "المدخلات", "صوت الطفل أو كتابته", "بالعربية أو الإنجليزية"),
             (";823;p24", "02", "المعالجة", "سياسة الدور والحراس", "واسترجاع من البنك"),
             (";824;p24", "03", "المخرجات", "رد قصير بصوت دافئ", "وبطاقة حين يسأل"),
             (";825;p24", "04", "الأثر", "قيمة يمارسها الطفل", "ومصدر يراه الوالدان")]
    for name, n, t, l1, l2 in steps:
        top_anchor(fill(s, by_name(s, name), [n, t, l1, l2]))
    caption(s, "الشيفرة: turn_pipeline.py مسار واحد للصوت والكتابة · retrieval.py · turn_policy.py · فرع hackathon", y=9.2)

    # A2 Key message (light) ---------------------------------------------------
    s = new(16)
    # this layout has no section label: add one where the other slides carry it
    add_text(s, 11.33, 1.12, 7.08, 0.46, ["ملحق · آلية العمل"], size=MIN_PT, color=PURPLE)
    fill(s, ph(s, 0), ["لا مصدر، لا جواب", "ولا فتوى، ولا نص يكتب من الذاكرة"])
    fill(s, ph(s, 1), ["إن لم يجد الصديق عنصرا مراجعا في البنك، يقول إنه لا يعرف مصدرا موثوقا ويحيل إلى الوالدين أو المعلم"])
    set_geom(ph(s, 1), x=2.71, w=14.58, h=1.2)
    fill(s, ph(s, 2), ["القرار رقم 3 في خطة المشروع PLAN.md، وتطبقه سياسة الدور في كل دور صوتا وكتابة"])

    # A3 Content rules (light four points) ------------------------------------
    s = new(20)
    kicker(s, "ملحق · آلية العمل")
    fill(s, ph(s, 0), ["منهج المحتوى والسلامة"])
    fill(s, ph(s, 2), ["قواعد ثابتة تطبقها سياسة الدور في كل رد، صوتا وكتابة"])
    rules = [(3, 4, "الرجاء أولا", "نبدأ برحمة الله ومحبته قبل الوعيد، بقرار من القائد بعد استشارة عالم"),
             (5, 6, "لا فتاوى", "المسائل الشخصية في الفقه تحال إلى الوالدين أو عالم، ولا يصدر الصديق حكما"),
             (7, 8, "العقيدة والخلاف", "مسائل العقيدة الحساسة تحال، وفي مسائل الخلاف يذكر المتفق عليه فقط"),
             (9, 13, "الطمأنة والأمان", "الحزن والقلق يقابلهما بالطمأنة، والخطر يحوله إلى بالغ موثوق")]
    for a, b, t, d in rules:
        fill(s, ph(s, a), [t])
        fill(s, ph(s, b), [d])
    caption(s, "المصدر: قرارات القائد في build-progress.md · PR #53 وPR #56 · turn_rules.json", y=9.55, dark=False)

    # A4 Content review workflow -------------------------------------------
    s = new(24)
    kicker(s, "ملحق · التشغيل")
    fill(s, ph(s, 0), ["سير مراجعة المحتوى"])
    fill(s, ph(s, 2), ["لا يصل إلى الطفل إلا عنصر راجعه إنسان، وأي تعديل على نصه يبطل المراجعة"])
    wf = [(";822;p24", "01", "النسخ من المصدر", "آليا من صفحة المصدر", "مع تحقق بالبصمة"),
          (";823;p24", "02", "مسودة للمراجعة", "شرح للطفل بلغتين", "لعمر 6 إلى 9 و10 إلى 13"),
          (";824;p24", "03", "مراجعة بشرية", "يوقعها مراجع مسمى", "وتسجل ببصمة العنصر"),
          (";825;p24", "04", "الخدمة للطفل", "لا يخدم إلا المراجع", "وتعديل النص يبطلها")]
    for name, n, t, l1, l2 in wf:
        top_anchor(fill(s, by_name(s, name), [n, t, l1, l2]))
    caption(s, "المراجعون: عبدالرحمن سلامة، وعبدالرحمن محملجي لأقسام الدرر · المصدر: build-progress.md وdeliverables/sources.md", y=9.2)

    # A5 Avatar: four state frames, Arabic labels over the frames' English ones --------
    s = new(14)
    kicker(s, "ملحق · النموذج الأولي")
    fill(s, ph(s, 0), ["شخصية تسمع وتفكر وتتكلم"])
    remove(ph(s, 2))
    th = 3.3
    tw = th * 270 / 318
    tgap = 0.24
    gx0 = 1.75 + (6.9 - (2 * tw + tgap)) / 2
    gy0 = 3.2
    labels = ["هدوء", "إنصات", "كلام", "كلام مع حركة يد"]
    for i, (path, lab) in enumerate(zip(imgs["avatar_frames"], labels)):
        col, row = 1 - (i % 2), i // 2          # right to left, then down
        tx, ty = gx0 + col * (tw + tgap), gy0 + row * (th + tgap)
        add_picture(s, path, tx, ty, tw, th, radius=0.05)
        # the frame's own label sits at 10..126 x 10..46 px of 270 x 318; cover it fully
        k = tw / 270
        lw = max(1.45, 0.24 + 0.155 * len(lab))
        add_box(s, tx + 5 * k, ty + 5 * k, lw, 0.47, fill=PILL, radius=0.45)
        add_text(s, tx + 5 * k, ty + 5 * k, lw, 0.47, [lab], size=MIN_PT, bold=True, color="FFFFFF",
                 align="c", anchor="m")
    fill(s, ph(s, 3), ["مزامنة الشفاه مع الصوت الحقيقي"])
    fill(s, ph(s, 4), ["يحلل المتصفح صوت الرد ليحرك الفم بأربعة عشر شكلا بالعربية والإنجليزية، وتتبدل الحالة بين الهدوء والإنصات والكلام"])
    fill(s, ph(s, 5), ["قبل التحدي كان الفك وحده يتحرك · 80 اختبارا للمزامنة",
                       "المصدر: handoffs/06-lipsync-ar.md وproduct-web.md", "5 أكتوبر 2026"])

    # A6 Accessibility (wide image) -------------------------------------------
    s = new(13)
    kicker(s, "ملحق · النموذج الأولي")
    fill(s, ph(s, 0), ["سهل للطفل، ومتاح بلغتين"])
    p = ph(s, 2)
    remove(p)
    add_picture(s, imgs["access"], 1.2, 3.45, 9.95, 9.95 * 5.16 / 9.17)
    fill(s, ph(s, 3), ["صوت أولا، وكتابة عند الحاجة"])
    fill(s, ph(s, 4), ["زر واحد للحديث، ووضع كتابة لمن لا يريد الكلام، وردود قصيرة لعمر 6 إلى 13 سنة"])
    fill(s, ph(s, 5), ["عربي أولا، ومتاح للجميع"])
    fill(s, ph(s, 6), ["من اليمين لليسار مع الإنجليزية، وشارة ذكاء اصطناعي دائمة، ووضع تقليل الحركة"])
    add_text(s, 11.44, 9.3, 6.98, 0.9, ["0 مخالفات في فحص axe على 12 صفحة باللغتين", "المصدر: product-web.md القسم 16 · 5 أكتوبر 2026"],
             size=MIN_PT, color=MUTED)

    # A7 Quality gate, before and after (light table) ----------------------------------
    # Both runs judged by Opus with the lead's rubric on the lead's sealed sets (610 turns, 20 conversations).
    # Before = first run, 5 Oct, d794bc7 (74 verses, before the card gate); after = re-run, 6 Oct, 8cbe64a
    # (273 items, card gate in; the same product code as 2c04305). All 11 rows, so the row that got worse
    # (declined) and the rows that did not move are on the slide too. The verdict is written in the cell,
    # not shown by colour alone. Numbers: eval-reports/quality-gate-summary.md sections 1 and 3.
    s = new(21)
    kicker(s, "ملحق · الأثر")
    fill(s, ph(s, 0), ["بوابة جودة المحادثة: قبل وبعد"])
    fill(s, ph(s, 2), ["قبل بوابة البطاقات: البناء d794bc7 و74 آية · بعدها: البناء 8cbe64a و273 عنصرا · 610 أدوار مختومة ومحكم Opus"])
    gf = [sh for sh in s.shapes if sh.has_table][0]
    gf.left, gf.top = I(1.54), I(3.75)
    ok = lambda t: {"text": t + " · ينجح", "bold": True, "color": PURPLE}
    no = lambda t: {"text": t + " · يفشل"}
    cols = [6.6, 2.9, 3.0, 4.38]
    fill_table(s, gf, [
        ["البوابة", "الهدف", "قبل · 5 أكتوبر", "بعد · 6 أكتوبر"],
        ["إقحام درس في الدردشة العادية", "10% أو أقل", "19.8%", ok("9.9%")],
        ["درس عند التحية أو الشكر أو الوداع", "0%", "59.1%", no("4.5%")],
        ["معلومات مختلقة، بالعدد", "0", "37", no("34")],
        ["إجابات مستندة إلى مصدر", "90% أو أكثر", "46.7%", no("68.9%")],
        ["اعتذار حين لا مصدر، دون كلام من الذاكرة", "95% أو أكثر", "89.3%", {"text": "78.6% · يفشل، وتراجع"}],
        ["إحالة دون فتوى", "100%", "66.7%", no("83.3%")],
        ["لا تأكيد ولا نفي لما ينقله الطفل", "100%", "83.3%", no("83.3%")],
        ["مسار الطمأنة والأمان", "100%", "63.6%", no("63.6%")],
        ["الهوية: ليس إنسانا ولا يسمي المزود ولا يتلو", "100%", "93.3%", ok("100%")],
        ["جودة الصديق في الدور، من 5", "4 أو أكثر", "3.61", no("3.96")],
        ["جودة الصديق في المحادثة، من 5", "4 أو أكثر", "2.35", no("3.20")],
    ], col_widths_rtl=cols, size=MIN_PT, header_size=20, row_heights=[0.48] + [0.42] * 11)
    src = ph(s, 3)
    set_geom(src, y=9.03, h=1.3)
    fill(s, src, ["الهلوسات 34 بقاعدة العد (نتيجة الشيفرة تغلب المحكم): 33 من المحكم وإنذار كاذب واحد من أداة القياس · مسار الطمأنة 81.8% لو عد بطريقة التشغيل الأول · أدوار الشخصيات بعد الافتتاح يكتبها المحاكي فتختلف بين التشغيلين",
                  "المصدر: quality-gate-summary.md القسمان 1 و3، الأرقام فقط لأن التقرير الكامل يقتبس جملا مختومة · لا تقارن بالمؤشر التقريبي في الشريحة التالية"])

    # A8 Lesson-creep proxy bars ---------------------------------------------------
    s = new(22)
    kicker(s, "ملحق · الأثر")
    fill(s, ph(s, 0), ["مؤشر القائد التقريبي: الدرس حين يطلبه الطفل فقط"])
    fill(s, ph(s, 2), ["نسبة الأدوار العادية التي أظهرت بطاقة، من قرارات سياسة الدور دون نموذج، قبل بوابة البطاقات (PR #56) وبعدها"])
    for sh in list(s.shapes):
        if sh.is_placeholder or sh.shape_type == 13:   # keep placeholders and pictures (bg, logo)
            continue
        remove(sh)
    base_y, scale = 8.05, 4.0 / 55.0
    groups = [(9.6, "مجموعة القائد", 51.5), (4.3, "مجموعته المحجوزة", 45.5)]
    bwid = 1.75
    for cx, gname, before in groups:
        for k, (lab, val, col) in enumerate([("قبل", before, PURPLE), ("بعد", 0.0, TURQ)]):
            bx = cx + 0.18 if k == 0 else cx - 0.18 - bwid
            hgt = max(val * scale, 0.07)
            add_box(s, bx, base_y - hgt, bwid, hgt, fill=col, radius=0.12 if val else 0.5)
            add_text(s, bx, base_y - hgt - 0.72, bwid, 0.62, [f"{val:g}%"], size=30, bold=True,
                     color=OFF, align="c", anchor="b")
            add_text(s, bx, base_y + 0.12, bwid, 0.45, [lab], size=20, color=BODY, align="c")
        add_text(s, cx - 2.2, base_y + 0.6, 4.4, 0.5, [gname], size=22, bold=True, color=OFF, align="c")
    ln = s.shapes.add_connector(1, I(1.9), I(base_y), I(12.3), I(base_y))
    _line_xml(ln._element.spPr, CARDLINE, 1.25)
    # facts card on the right
    add_box(s, 13.25, 3.75, 5.17, 4.95, fill="FFFFFF", fill_alpha=4, line=TURQ, line_alpha=15)
    add_text(s, 13.6, 3.95, 4.5, 0.6, ["البطاقة في موضعها"], size=24, bold=True, color=TURQ)
    add_text(s, 13.6, 4.75, 4.5, 1.7,
             [{"text": "0 من 55", "size": 40, "bold": True, "color": OFF},
              {"text": "عبارة لطف أو تحية أظهرت بطاقة", "size": 20, "color": BODY}])
    add_text(s, 13.6, 6.6, 4.5, 1.7,
             [{"text": "40 من 40", "size": 40, "bold": True, "color": OFF},
              {"text": "سؤالا يستحق مصدرا ما زال يناله", "size": 20, "color": BODY}])
    src = ph(s, 3)
    set_geom(src, x=1.54, w=16.88, y=9.35, h=0.9)
    fill(s, src, ["المصدر: quality-gate-summary.md القسم 2، منقولا عن وصف PR #56 في 5 أكتوبر · مجموعات القائد ليست في المستودع · مؤشر تقريبي وليس نسبة محكمة، فلا يقارن بالبوابة المحكمة في الشريحة السابقة"])

    # A-cost-1 The USD estimate per session: measured units x public prices (light table) -------------
    it = EST["items"]
    s = new(21)
    kicker(s, "ملحق · التشغيل")
    fill(s, ph(s, 0), ["تقدير التكلفة: الجلسة الواحدة"])
    fill(s, ph(s, 2), ["تقدير لا قياس: وحدات مقيسة في المستودع، أو مفترضة حيث لم تقس، مضروبة في أسعار المزودين المعلنة"])
    gf = [sh for sh in s.shapes if sh.has_table][0]
    gf.left, gf.top = I(1.54), I(3.75)
    cols = [3.3, 6.9, 4.3, 2.38]
    rng = lambda k: f"{it[k][0]} إلى {it[k][1]}"
    fill_table(s, gf, [
        ["البند", "الوحدات لكل جلسة", "السعر المعلن", "للجلسة، دولار"],
        ["نموذج المحادثة\ngpt-5.4-mini",
         "4 استدعاءات (مقيس) إلى 12 (افتراض). للاستدعاء 6.5 ألف رمز إدخال و109 إخراج، متوسط تشغيل "
         "التقييم مع المحكم، حدا أعلى؛ و8 آلاف إدخال للجلسة الكاملة (افتراض)",
         "0.75 دولار لمليون رمز إدخال، و4.50 لمليون إخراج", rng("agent LLM")],
        ["تقارير ما بعد الجلسة\nبالنموذج نفسه",
         "4 استدعاءات (مقيس)، لكل استدعاء ألف إلى 6 آلاف رمز إدخال و200 إلى 800 إخراج (افتراض)",
         "السعر نفسه", rng("reporting LLM")],
        ["الكلام إلى نص\ngpt-4o-mini-transcribe",
         "1.5 إلى 5 دقائق، أي طول الجلسة، لأن النسخ الفوري مفعل افتراضيا ويبقى مفتوحا طوالها (افتراض)",
         "نحو 0.003 دولار للدقيقة، بتقدير OpenAI", rng("STT")],
        ["الصوت\nElevenLabs Flash v2.5",
         "300 حرف لتحية وسؤالين (مقيس) إلى 2500 حرف لجلسة 5 دقائق (تقدير البروفة)",
         "0.04 دولار لكل ألف حرف، في كل الباقات", rng("TTS")],
        ["LiveKit Cloud", "طرفان، الطفل والوكيل، لمدة 1.5 إلى 5 دقائق (افتراض)",
         "0.0005 دولار للمشارك في الدقيقة، و0.01 للوكيل في الحد الأعلى", rng("LiveKit")],
        [{"text": "المجموع للجلسة", "bold": True},
         {"text": "من جلسة قصيرة إلى جلسة كاملة مدتها 5 دقائق (حد العرض)"},
         {"text": ""}, {"text": EST["session"], "bold": True, "color": PURPLE}],
    ], col_widths_rtl=cols, size=MIN_PT, header_size=20,
        row_heights=[0.55, 1.1, 0.8, 0.8, 0.8, 0.8, 0.55])
    src = ph(s, 3)
    set_geom(src, x=1.54, w=16.88, y=9.4, h=0.95)
    fill(s, src, [[{"text": "الأسعار كما فحصت في 6 أكتوبر 2026: "},
                   {"text": "platform.openai.com/docs/pricing", "link": "https://platform.openai.com/docs/pricing"},
                   {"text": " · "},
                   {"text": "elevenlabs.io/pricing/api", "link": "https://elevenlabs.io/pricing/api"},
                   {"text": " · "},
                   {"text": "livekit.com/pricing", "link": "https://livekit.com/pricing"}],
                  "الوحدات المقيسة: latency/LOAD-REHEARSAL.md وeval-reports/lead-summary.json · الحساب: deck-src/cost_estimate.py"])

    # A-cost-2 The USD estimate per month, and the assumptions (two options) ------------------------
    s = new(28)
    kicker(s, "ملحق · التشغيل")
    fill(s, ph(s, 0), ["تقدير التكلفة: الشهر والافتراضات"])
    fill(s, ph(s, 2), ["تقدير لا قياس: تكلفة الجلسة من الشريحة السابقة مضروبة في عدد جلسات مفترض"])
    js, fs = EST["judge_sessions"], EST["family_sessions"]
    fill(s, ph(s, 3), ["30 محكما"])
    fill(s, ph(s, 4), ["طوال فترة التحكيم، من 7 إلى 22 أكتوبر"])
    fill(s, ph(s, 5), ["100 أسرة"])
    fill(s, ph(s, 6), ["طفل واحد لكل أسرة، في الشهر"])
    gap_p = {"text": "", "size": 8}   # the template's blank spacer paragraphs, kept small
    fill(s, ph(s, 7), ["2 إلى 4 جلسات لكل محكم (افتراض)", gap_p, f"{js[0]} إلى {js[1]} جلسة", gap_p,
                       {"text": EST["judges"] + " دولارا للفترة كلها", "bold": True, "color": TURQ, "size": 25.5}],
         size=22)
    fill(s, ph(s, 8), ["12 إلى 20 جلسة لكل طفل في الشهر (افتراض)", gap_p, f"{fs[0]:,} إلى {fs[1]:,} جلسة", gap_p,
                       {"text": EST["families"] + " دولارا في الشهر", "bold": True, "color": TURQ, "size": 25.5}],
         size=22)
    add_text(s, 1.54, 8.3, 16.88, 2.1,
             [{"text": "الافتراضات، وما لم يدخل في التقدير", "bold": True, "color": TURQ, "size": 20},
              "الجلسة القصيرة 1.5 دقيقة، والكاملة 5 دقائق بنحو 10 أسئلة، وهي حد العرض · قد تطول جلسات الأسر، والتقدير يفترض الحد نفسه",
              "رموز النموذج: متوسط تشغيل التقييم مع المحكم، حدا أعلى للاستدعاء · رموز التقارير لم تقس · لا خصم على المدخلات المخزنة مؤقتا",
              "نفترض أن LiveKit يحسب وكيلنا مشاركا لأنه يعمل على خادمنا، والحد الأعلى يضيف دقائق جلسة الوكيل احتياطا",
              "خارج التقدير: الخادم، ورسوم الباقات: LiveKit Build مجانية وShip من 50 دولارا شهريا، وElevenLabs من 6 دولارات"],
             size=MIN_PT, color=BODY, spacing=3)

    # A9 Built vs existed (two options) ---------------------------------------
    s = new(28)
    kicker(s, "ملحق · الإفصاح")
    fill(s, ph(s, 0), ["ما كان موجودا، وما بنيناه في التحدي"])
    fill(s, ph(s, 2), ["الإفصاح عن العمل السابق كما يطلب دليل المشاركين، وكل تغيير في سجل git بتاريخه"])
    fill(s, ph(s, 3), ["قبل التحدي"])
    fill(s, ph(s, 4), ["مشروع تخرج سابق، معلن في DISCLOSURE.md"])
    fill(s, ph(s, 5), ["خلال 4 إلى 6 أكتوبر"])
    fill(s, ph(s, 6), ["العمل الذي يحكم عليه"])
    fill(s, ph(s, 7), ["تطبيق ويب وخادم Django، والجوال قيد التطوير", "", "وكيل صوتي بصوت xAI، وشخصية ثلاثية الأبعاد", "",
                       "6 موضوعات أخلاقية و19 مرجعا نصيا بلا رابط"])
    fill(s, ph(s, 8), ["بنك 273 عنصرا مراجعا و38 قيمة", "", "سياسة الدور والحراس وبطاقات المصادر", "",
                       "صوت ElevenLabs ومزامنة الشفاه والتقييم"])

    # A10 Roadmap timeline ----------------------------------------------------
    s = new(18)
    kicker(s, "ملحق · خارطة الطريق")
    fill(s, ph(s, 0), ["خارطة الطريق"])
    fill(s, ph(s, 2), ["ما أنجز، وما نقترحه بعد التحدي"])
    tl = [(3, 4, 5, 6, "01", "مشروع التخرج", "قبل 4 أكتوبر", "تطبيق لقيمة الصدق بالصوت والشخصية"),
          (7, 8, 9, 13, "02", "التحدي", "4 إلى 6 أكتوبر", "273 عنصرا مراجعا، وبطاقات، وصوت جديد"),
          (14, 15, 16, 17, "03", "التحكيم", "حتى 22 أكتوبر", "عرض حي مستقر، ومراقبة للإنفاق"),
          (18, 19, 20, 21, "04", "تجربة الأسر", "مقترح", "تجربة مع أسر وشريك مراجعة علمي"),
          (22, 23, 24, 25, "05", "التوسع", "مقترح", "قيم وتلاوات أكثر، وتطبيق جوال قيد التطوير")]
    for a, b, c, d, n, st, per, outp in tl:
        fill(s, ph(s, a), [n])
        fill(s, ph(s, b), [st])
        fill(s, ph(s, c), [per])
        fill(s, ph(s, d), [outp])

    # ---------------------------------------------------------- remove guide + samples
    lst = prs.slides._sldIdLst
    for sid in orig_ids:
        prs.part.drop_rel(sid.rId)
        lst.remove(sid)
    # The dropped template slides are no longer reachable, so they are not saved. Name the kept
    # slide parts slide1..slideN in show order, so the package does not start at slide32.xml and
    # look as if 31 template slides were still inside it.
    prs.part.rename_slide_parts([sid.rId for sid in lst])

    # ---------------------------------------------------------- remove unused layouts
    used = {id(sl.slide_layout.part) for sl in prs.slides}
    for master in prs.slide_masters:
        lid_lst = master._element.find(qn("p:sldLayoutIdLst"))
        for lid in list(lid_lst):
            rid = lid.get(qn("r:id"))
            part = master.part.related_part(rid)
            if id(part) not in used:
                master.part.drop_rel(rid)
                lid_lst.remove(lid)

    # ---------------------------------------------------------- empty placeholders
    for sl in prs.slides:
        for sh in list(sl.shapes):
            if sh.is_placeholder and sh.placeholder_format.type not in (13,):   # keep slide-number fields
                if sh.has_text_frame and not sh.text_frame.text.strip():
                    remove(sh)

    set_theme_links(prs)
    prs.core_properties.title = "الصديق الصدوق"
    prs.core_properties.subject = "تحدي الذكاء الاصطناعي في خدمة المحتوى الإسلامي"
    prs.save(str(OUTFILE))
    return OUTFILE


# Template chrome, not our text: the organisers' footer line (13.5 pt) and the closing slide's
# partners label, whose colon is a separate run that inherits its 18 pt size.
CHROME = {"تحدي الذكاء الاصطناعي في خدمة المحتوى الإسلامي", ":"}


def gate(path):
    """Fail on template brackets, missing fonts, leftover guide text, or any of our runs under MIN_PT."""
    problems = []
    with zipfile.ZipFile(path) as z:
        slides = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                        key=lambda n: int(re.findall(r"\d+", n)[0]))
        for n in slides:
            x = z.read(n).decode("utf8")
            root = etree.fromstring(x.encode("utf8"))
            for r in root.iter(qn("a:r")):
                t = (r.findtext(qn("a:t")) or "")
                if "[" in t or "]" in t:
                    problems.append(f"{n}: bracket in text: {t[:60]}")
                if "دليل الاستخدام" in t or "قالب عروض" in t:
                    problems.append(f"{n}: guide text: {t[:60]}")
                rpr = r.find(qn("a:rPr"))
                sz = rpr.get("sz") if rpr is not None else None
                if t.strip() and t.strip() not in CHROME:
                    if sz is None:
                        problems.append(f"{n}: run with no explicit size: {t[:40]}")
                    elif int(sz) < MIN_PT * 100:
                        problems.append(f"{n}: {int(sz) / 100:g} pt run: {t[:40]}")
            if "fontScale" in x:
                problems.append(f"{n}: autofit fontScale would shrink text")
            for face in set(re.findall(r'typeface="([^"]+)"', x)):
                if not face.startswith("Readex Pro") and face not in ("Arial",):
                    problems.append(f"{n}: font {face}")
        print("slides:", len(slides))
    return problems


if __name__ == "__main__":
    out = build()
    probs = gate(out)
    print("built", out, out.stat().st_size // 1024, "KB")
    if probs:
        print("GATE PROBLEMS:")
        for p in probs:
            print("  ", p)
        sys.exit(1)
    print("gate: OK")
